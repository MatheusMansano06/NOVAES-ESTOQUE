import openpyxl
import pytest

from app.negociacao_shopee import (
    COL,
    PlanilhaInvalida,
    ShopeeIndisponivel,
    coletar_precos,
    escrever,
    ler_planilha,
    montar,
    resolver,
)

CABECALHO = {
    1: "Flag", 4: "Cluster", 5: "Categoria L1", 6: "SKU", 7: "Preço Referência",
    8: "ADGMV L30D", 9: "ADO L30D", 10: "Rebate Médio L30D", 11: "Score",
    14: "Item ID", 16: "Model ID", 17: "Model Name", 18: "Estoque D-1",
    19: "Estoque Full D-1", 20: "Tipo Negociação", 21: "Preço Original",
    22: "Preço Site D-1", 23: "Preço do Seller", 30: "Preço (R$)",
    31: "Estoque Negociado",
}

# (item_id, model_id, sku, descricao, preco_ref, estoque_d1, full_d1, preco_site)
PRODUTOS = [
    ("23893446809", "169717255440", "SKU-A", "Protetor escapamento", 70.79, 608, 12, 67.79),
    ("21598020157", "228778685523", "SKU-B", "Protetor Titan 160", 61.72, 285, 294, 78.99),
    ("22098064287", "199601311893", "SKU-C", "Par retrovisor", 20.00, 14, 48, 19.90),
]


def planilha(tmp_path, cabecalho=None, nome_aba="Pontual - Nível Model"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = nome_aba
    for coluna, rotulo in (cabecalho or CABECALHO).items():
        ws.cell(2, coluna).value = rotulo
    for i, p in enumerate(PRODUTOS):
        linha = 3 + i
        ws.cell(linha, COL["item_id"]).value = int(p[0])
        ws.cell(linha, COL["model_id"]).value = int(p[1])
        ws.cell(linha, COL["sku"]).value = p[2]
        ws.cell(linha, COL["descricao"]).value = p[3]
        ws.cell(linha, COL["preco_referencia"]).value = p[4]
        ws.cell(linha, COL["estoque_d1"]).value = p[5]
        ws.cell(linha, COL["estoque_full_d1"]).value = p[6]
        ws.cell(linha, COL["preco_site_d1"]).value = p[7]
        ws.cell(linha, COL["score"]).value = 900 + i
        ws.cell(linha, 32).value = "=AD{0}*AE{0}".format(linha)   # AF: fórmula vizinha
    caminho = str(tmp_path / "negociacao.xlsx")
    wb.save(caminho)
    wb.close()
    return caminho


PRECOS = {
    "169717255440": {"preco": 79.99, "campanha_id": "1", "campanha_nome": "SETEMBRO", "campanhas": 1},
    "21598020157": {"preco": 79.99, "campanha_id": "2", "campanha_nome": "AGOSTO", "campanhas": 2},
}
ESTOQUES = {
    "169717255440": {"seller": 767, "shopee": 0},
    "21598020157": {"seller": 0, "shopee": 289},
}


def test_le_as_colunas_certas(tmp_path):
    linhas = ler_planilha(planilha(tmp_path))
    assert len(linhas) == 3
    primeira = linhas[0]
    assert primeira["linha"] == 3
    assert primeira["item_id"] == "23893446809"
    assert primeira["model_id"] == "169717255440"
    assert primeira["sku"] == "SKU-A"
    assert primeira["preco_referencia"] == 70.79
    assert primeira["estoque_d1"] == 608
    assert primeira["estoque_full_d1"] == 12
    assert primeira["extras"]["score"] == 900


def test_recusa_planilha_com_layout_diferente(tmp_path):
    trocado = dict(CABECALHO)
    trocado[30], trocado[31] = "Estoque Negociado", "Preço (R$)"
    with pytest.raises(PlanilhaInvalida, match="layout"):
        ler_planilha(planilha(tmp_path, cabecalho=trocado))


def test_recusa_arquivo_sem_a_aba(tmp_path):
    with pytest.raises(PlanilhaInvalida, match="aba"):
        ler_planilha(planilha(tmp_path, nome_aba="Resumo"))


def test_resolver_prefere_model_e_cai_no_item():
    mapa = {"169717255440": "por-model", "21598020157": "por-item"}
    assert resolver(mapa, "23893446809", "169717255440") == "por-model"
    assert resolver(mapa, "21598020157", "228778685523") == "por-item"
    assert resolver(mapa, "999", "888") is None


def test_montar_conta_zerados_multicampanha_e_faltantes(tmp_path):
    resultado = montar(ler_planilha(planilha(tmp_path)), PRECOS, ESTOQUES)
    assert resultado["total_linhas"] == 3
    assert resultado["total_zerados"] == 1          # SKU-B, seller=0
    assert resultado["total_multi_campanha"] == 1   # SKU-B em duas campanhas
    assert resultado["sem_preco"] == ["22098064287"]
    assert resultado["sem_estoque"] == ["22098064287"]

    por_item = {i["item_id"]: i for i in resultado["itens"]}
    assert por_item["23893446809"]["preco_preenchido"] == 79.99
    assert por_item["23893446809"]["estoque_preenchido"] == 767
    assert por_item["23893446809"]["estoque_full"] == 0
    assert por_item["21598020157"]["campanha_nome"] == "AGOSTO"
    assert por_item["22098064287"]["preco_preenchido"] is None


def test_escrever_preenche_ad_ae_e_nao_quebra_formula(tmp_path):
    origem = planilha(tmp_path)
    destino = str(tmp_path / "preenchida.xlsx")
    resultado = montar(ler_planilha(origem), PRECOS, ESTOQUES)

    assert escrever(origem, destino, resultado["itens"]) == 2

    wb = openpyxl.load_workbook(destino)
    ws = wb["Pontual - Nível Model"]
    assert ws.cell(3, COL["preco"]).value == 79.99
    assert ws.cell(3, COL["estoque"]).value == 767
    assert ws.cell(4, COL["estoque"]).value == 0          # zero é valor, não vazio
    assert ws.cell(5, COL["preco"]).value is None          # sem dado, fica em branco
    assert ws.cell(3, 32).value == "=AD3*AE3"              # fórmula vizinha intacta
    wb.close()


class ClienteFake:
    def __init__(self, promocoes, precos_por_campanha):
        self._promocoes = promocoes
        self._precos = precos_por_campanha

    def listar_promocoes(self, status="ongoing"):
        return {"promocoes": self._promocoes}

    def precos_da_promocao(self, discount_id):
        return {"precos": self._precos[discount_id]}


def test_coletar_precos_ignora_campanha_que_nao_e_do_seller():
    cliente = ClienteFake(
        [{"discount_id": 1, "nome": "SETEMBRO", "origem": 0},
         {"discount_id": 9, "nome": "CAMPANHA SHOPEE", "origem": 1}],
        {1: {"m1": 50.0}, 9: {"m1": 10.0}},
    )
    precos = coletar_precos(cliente)
    assert precos["m1"]["preco"] == 50.0
    assert precos["m1"]["campanhas"] == 1


def test_coletar_precos_usa_o_menor_entre_campanhas_proprias():
    cliente = ClienteFake(
        [{"discount_id": 1, "nome": "SETEMBRO", "origem": 0},
         {"discount_id": 2, "nome": "AGOSTO", "origem": 0}],
        {1: {"m1": 79.99}, 2: {"m1": 69.99}},
    )
    precos = coletar_precos(cliente)
    assert precos["m1"]["preco"] == 69.99
    assert precos["m1"]["campanha_nome"] == "AGOSTO"
    assert precos["m1"]["campanhas"] == 2


def test_coletar_precos_sem_campanha_propria_vira_indisponivel():
    cliente = ClienteFake([{"discount_id": 9, "nome": "SHOPEE", "origem": 1}], {})
    with pytest.raises(ShopeeIndisponivel):
        coletar_precos(cliente)
