"""Fluxo completo da aba: sobe planilha, preenche, baixa, histórico e BI.
A Shopee entra por monkeypatch — o que se testa aqui é o encanamento."""

import io

import openpyxl
import pytest
from starlette.applications import Starlette
from starlette.routing import Route
from starlette.testclient import TestClient

from app import main
from app import negociacao_shopee as negoc
from app.models import NegociacaoShopee, NegociacaoShopeeItem
from database import SessionLocal

from tests.test_negociacao_shopee import CABECALHO, PRODUTOS

app_teste = Starlette(routes=[
    Route("/api/negociacoes-shopee/bi", main.negoc_bi, methods=["GET"]),
    Route("/api/negociacoes-shopee", main.negoc_listar, methods=["GET"]),
    Route("/api/negociacoes-shopee", main.negoc_criar, methods=["POST"]),
    Route("/api/negociacoes-shopee/{id:int}", main.negoc_detalhe, methods=["GET"]),
    Route("/api/negociacoes-shopee/{id:int}/arquivo/{idx:int}", main.negoc_arquivo, methods=["GET"]),
    Route("/api/negociacoes-shopee/{id:int}/reprocessar", main.negoc_reprocessar, methods=["POST"]),
])
cliente = TestClient(app_teste)

PRECOS = {
    "169717255440": {"preco": 79.99, "campanha_id": "1", "campanha_nome": "SETEMBRO", "campanhas": 1},
    "21598020157": {"preco": 59.99, "campanha_id": "2", "campanha_nome": "AGOSTO", "campanhas": 1},
    "199601311893": {"preco": 17.99, "campanha_id": "1", "campanha_nome": "SETEMBRO", "campanhas": 1},
}
ESTOQUES = {
    "169717255440": {"seller": 767, "shopee": 0},
    "21598020157": {"seller": 0, "shopee": 289},
    "199601311893": {"seller": 405, "shopee": 6},
}


@pytest.fixture(autouse=True)
def limpar_banco():
    db = SessionLocal()
    db.query(NegociacaoShopeeItem).delete()
    db.query(NegociacaoShopee).delete()
    db.commit()
    db.close()


@pytest.fixture
def shopee_ok(monkeypatch):
    monkeypatch.setattr(negoc, "coletar_precos", lambda cliente: PRECOS)
    monkeypatch.setattr(negoc, "coletar_estoques", lambda cliente, ids: ESTOQUES)


def bytes_planilha(produtos=PRODUTOS):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Pontual - Nível Model"
    for coluna, rotulo in CABECALHO.items():
        ws.cell(2, coluna).value = rotulo
    for i, p in enumerate(produtos):
        linha = 3 + i
        ws.cell(linha, negoc.COL["item_id"]).value = int(p[0])
        ws.cell(linha, negoc.COL["model_id"]).value = int(p[1])
        ws.cell(linha, negoc.COL["sku"]).value = p[2]
        ws.cell(linha, negoc.COL["descricao"]).value = p[3]
        ws.cell(linha, negoc.COL["preco_referencia"]).value = p[4]
        ws.cell(linha, negoc.COL["estoque_d1"]).value = p[5]
        ws.cell(linha, negoc.COL["estoque_full_d1"]).value = p[6]
        ws.cell(linha, negoc.COL["preco_site_d1"]).value = p[7]
    buffer = io.BytesIO()
    wb.save(buffer)
    wb.close()
    return buffer.getvalue()


def subir(nomes=("Outubro.xlsx",), produtos=PRODUTOS, **campos):
    arquivos = [("arquivos", (n, bytes_planilha(produtos), "application/vnd.ms-excel"))
                for n in nomes]
    return cliente.post("/api/negociacoes-shopee",
                        files=arquivos,
                        data={"nome": "Outubro", "competencia": "2026-10", **campos})


def test_upload_preenche_e_grava_historico(shopee_ok):
    resp = subir()
    assert resp.status_code == 200
    corpo = resp.json()
    assert corpo["status"] == "preenchida"
    assert corpo["total_linhas"] == 3
    assert corpo["linhas_gravadas"] == 3
    assert corpo["total_zerados"] == 1
    assert corpo["sem_preco"] == []

    detalhe = cliente.get(f"/api/negociacoes-shopee/{corpo['id']}").json()
    por_item = {i["item_id"]: i for i in detalhe["itens"]}
    assert por_item["23893446809"]["preco"] == 79.99
    assert por_item["23893446809"]["estoque"] == 767
    assert por_item["21598020157"]["campanha"] == "AGOSTO"
    assert por_item["21598020157"]["estoque"] == 0
    # o dado que só a Shopee tem continua guardado
    assert por_item["23893446809"]["preco_referencia"] == 70.79
    assert por_item["23893446809"]["estoque_d1"] == 608


def test_dois_arquivos_viram_uma_negociacao_sem_duplicar_linha(shopee_ok):
    corpo = subir(nomes=("Outubro.xlsx", "spike 10.10.xlsx")).json()
    assert corpo["total_linhas"] == 3          # não 6
    assert corpo["linhas_gravadas"] == 6       # mas os dois arquivos foram escritos
    assert len(cliente.get("/api/negociacoes-shopee").json()[0]["arquivos"]) == 2


def test_download_devolve_xlsx_preenchido(shopee_ok):
    neg_id = subir().json()["id"]
    resp = cliente.get(f"/api/negociacoes-shopee/{neg_id}/arquivo/0")
    assert resp.status_code == 200
    assert "PREENCHIDO" in resp.headers["content-disposition"]

    ws = openpyxl.load_workbook(io.BytesIO(resp.content))["Pontual - Nível Model"]
    assert ws.cell(3, negoc.COL["preco"]).value == 79.99
    assert ws.cell(3, negoc.COL["estoque"]).value == 767
    assert ws.cell(4, negoc.COL["estoque"]).value == 0


def test_arquivo_inexistente_da_404(shopee_ok):
    neg_id = subir().json()["id"]
    assert cliente.get(f"/api/negociacoes-shopee/{neg_id}/arquivo/7").status_code == 404
    assert cliente.get("/api/negociacoes-shopee/9999").status_code == 404


def test_planilha_de_outro_formato_e_recusada_sem_sujar_o_historico(shopee_ok):
    wb = openpyxl.Workbook()
    wb.active.title = "Relatorio"
    buffer = io.BytesIO()
    wb.save(buffer)

    resp = cliente.post("/api/negociacoes-shopee",
                        files=[("arquivos", ("errado.xlsx", buffer.getvalue(), "application/vnd.ms-excel"))],
                        data={"nome": "Errado", "competencia": "2026-10"})
    assert resp.status_code == 400
    assert "aba" in resp.json()["erro"]
    assert cliente.get("/api/negociacoes-shopee").json() == []


def test_upload_sem_arquivo_da_400():
    resp = cliente.post("/api/negociacoes-shopee", data={"nome": "Vazio"})
    assert resp.status_code == 400


def test_shopee_fora_do_ar_deixa_pendente_e_reprocessa_depois(monkeypatch):
    def cai(cliente_shopee):
        raise negoc.ShopeeIndisponivel("token expirado")

    monkeypatch.setattr(negoc, "coletar_precos", cai)
    corpo = subir().json()
    assert corpo["status"] == "pendente"
    assert "token expirado" in corpo["erro"]

    # o arquivo cru já está guardado e baixável
    assert cliente.get(f"/api/negociacoes-shopee/{corpo['id']}/arquivo/0").status_code == 200

    monkeypatch.setattr(negoc, "coletar_precos", lambda c: PRECOS)
    monkeypatch.setattr(negoc, "coletar_estoques", lambda c, ids: ESTOQUES)
    resp = cliente.post(f"/api/negociacoes-shopee/{corpo['id']}/reprocessar").json()
    assert resp["status"] == "preenchida"
    assert resp["total_linhas"] == 3


def test_bi_traz_ruptura_preco_e_giro(shopee_ok):
    subir(nomes=("Setembro.xlsx",), produtos=PRODUTOS[:2], nome="Setembro", competencia="2026-09")
    subir(nomes=("Outubro.xlsx",), produtos=PRODUTOS[1:], nome="Outubro", competencia="2026-10")

    bi = cliente.get("/api/negociacoes-shopee/bi").json()
    assert bi["vazio"] is False
    assert bi["negociacao"]["nome"] == "Outubro"

    assert [r["item_id"] for r in bi["ruptura"]] == ["21598020157"]
    assert bi["ruptura"][0]["estoque_full"] == 289

    desvios = {p["item_id"]: p["desvio_pct"] for p in bi["preco"]}
    assert desvios["21598020157"] == pytest.approx(-2.8, abs=0.1)   # 59.99 vs 61.72
    assert desvios["22098064287"] == pytest.approx(-10.1, abs=0.1)  # 17.99 vs 20.00

    assert [e["item_id"] for e in bi["giro"]["entraram"]] == ["22098064287"]
    assert [s["item_id"] for s in bi["giro"]["sairam"]] == ["23893446809"]


def test_bi_compara_com_o_preco_de_site_quando_nao_ha_referencia(shopee_ok):
    """A Shopee só informa Preço Referência numa minoria das linhas — no resto
    a base tem de ser o Preço Site D-1, senão a visão de preço fica vazia."""
    sem_referencia = [(p[0], p[1], p[2], p[3], None, p[5], p[6], p[7]) for p in PRODUTOS]
    subir(produtos=sem_referencia)

    preco = {p["item_id"]: p for p in cliente.get("/api/negociacoes-shopee/bi").json()["preco"]}
    assert len(preco) == 3                                    # nenhuma linha se perdeu
    assert all(p["base_tipo"] == "site" for p in preco.values())
    # 79.99 contra os 67.79 que o produto tinha no site
    assert preco["23893446809"]["desvio_pct"] == pytest.approx(18.0, abs=0.1)
    assert preco["23893446809"]["referencia"] is None


def test_bi_sem_negociacao_nao_quebra():
    assert cliente.get("/api/negociacoes-shopee/bi").json() == {"vazio": True}
